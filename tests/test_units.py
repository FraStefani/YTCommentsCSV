"""Test offline (nessuna chiamata di rete): parsing, sanitizzazione e scrittura xlsx."""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

import xlsx_report
from xlsx_report import (
    TRUNCATE_MARK,
    build_output_path,
    clean_text,
    format_int,
    sanitize_filename,
    to_excel_datetime,
    write_report,
)
from yt_api import Comment, InvalidUrlError, VideoInfo, _comment_from, parse_video_id


class TestParseVideoId(unittest.TestCase):
    def test_formati_supportati(self) -> None:
        casi = {
            "https://youtu.be/ByZrTRph2-4?si=SuKq_vqCI-q9MRt7": "ByZrTRph2-4",
            "https://www.youtube.com/watch?v=ByZrTRph2-4&t=42s": "ByZrTRph2-4",
            "https://www.youtube.com/watch?list=PL123&v=ByZrTRph2-4": "ByZrTRph2-4",
            "https://www.youtube.com/shorts/ByZrTRph2-4": "ByZrTRph2-4",
            "https://www.youtube.com/live/ByZrTRph2-4": "ByZrTRph2-4",
            "https://www.youtube.com/embed/ByZrTRph2-4": "ByZrTRph2-4",
            "https://m.youtube.com/watch?v=ByZrTRph2-4": "ByZrTRph2-4",
            "youtu.be/ByZrTRph2-4": "ByZrTRph2-4",
            "ByZrTRph2-4": "ByZrTRph2-4",
            "  https://youtu.be/ByZrTRph2-4  ": "ByZrTRph2-4",
        }
        for url, expected in casi.items():
            with self.subTest(url=url):
                self.assertEqual(parse_video_id(url), expected)

    def test_url_non_validi(self) -> None:
        for url in ("", "https://vimeo.com/12345", "https://youtu.be/id-troppo-lungo-per-essere-valido",
                    "https://www.youtube.com/watch?v=non-valido!!", "https://www.youtube.com/"):
            with self.subTest(url=url):
                with self.assertRaises(InvalidUrlError):
                    parse_video_id(url)


class TestCommentMapping(unittest.TestCase):
    def test_campi_mancanti_non_alzano_errori(self) -> None:
        comment = _comment_from({}, is_reply=False)
        self.assertEqual(comment.author, "[utente rimosso]")
        self.assertEqual(comment.text, "")
        self.assertIsNone(comment.published_at)
        self.assertEqual(comment.like_count, 0)

    def test_entita_html_e_timestamp(self) -> None:
        payload = {
            "id": "Ugx123",
            "snippet": {
                "authorDisplayName": "Mario",
                "textDisplay": "Tom &amp; Jerry &#39;90",
                "publishedAt": "2024-03-14T18:30:00Z",
                "likeCount": 7,
            },
        }
        comment = _comment_from(payload, is_reply=True, parent_author="Luca")
        self.assertEqual(comment.text, "Tom & Jerry '90")
        self.assertEqual(comment.published_at, datetime(2024, 3, 14, 18, 30, tzinfo=timezone.utc))
        self.assertTrue(comment.is_reply)
        self.assertEqual(comment.parent_author, "Luca")


class TestSanitizzazione(unittest.TestCase):
    def test_nome_file(self) -> None:
        self.assertEqual(sanitize_filename('Canale/con:caratteri*vietati?'),
                         "Canale_con_caratteri_vietati_")
        self.assertEqual(sanitize_filename("   "), "canale_sconosciuto")
        self.assertEqual(sanitize_filename("nome."), "nome")
        self.assertEqual(sanitize_filename("CON"), "_CON")
        self.assertEqual(sanitize_filename("NUL.qualcosa"), "_NUL.qualcosa")
        self.assertLessEqual(len(sanitize_filename("x" * 400)), 150)

    def test_clean_text(self) -> None:
        self.assertEqual(clean_text("riga1\r\nriga2"), "riga1\nriga2")
        self.assertEqual(clean_text("con\x07campanella"), "concampanella")
        self.assertEqual(clean_text(None), "")
        self.assertEqual(clean_text("tab\tok"), "tab\tok")

    def test_troncamento(self) -> None:
        lungo = "a" * 40_000
        risultato = clean_text(lungo)
        self.assertEqual(len(risultato), xlsx_report.TEXT_LIMIT)
        self.assertTrue(risultato.endswith(TRUNCATE_MARK))

    def test_datetime_senza_timezone(self) -> None:
        aware = datetime(2024, 3, 14, 18, 30, tzinfo=timezone.utc)
        convertito = to_excel_datetime(aware, utc=True)
        self.assertIsNone(convertito.tzinfo)
        self.assertEqual(convertito, datetime(2024, 3, 14, 18, 30))
        self.assertIsNone(to_excel_datetime(None))

    def test_formattazione_numeri(self) -> None:
        self.assertEqual(format_int(1234567), "1.234.567")
        self.assertEqual(format_int(None), "N/D")

    def test_nome_progressivo_se_esiste(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            cartella = Path(tmp)
            primo = build_output_path("Canale", datetime(2026, 9, 23), cartella)
            self.assertEqual(primo.name, "Canale_2026-09-23.xlsx")
            primo.touch()
            secondo = build_output_path("Canale", datetime(2026, 9, 23), cartella)
            self.assertEqual(secondo.name, "Canale_2026-09-23_2.xlsx")


def _dati_di_prova() -> tuple[VideoInfo, list[Comment]]:
    base = datetime(2024, 3, 14, 18, 30, tzinfo=timezone.utc)
    video = VideoInfo(
        video_id="ByZrTRph2-4",
        url="https://www.youtube.com/watch?v=ByZrTRph2-4",
        title="Titolo di prova",
        channel="Canale di Prova",
        published_at=base,
        view_count=1_234_567,
        comment_count=3,
    )
    comments = [
        Comment("mario", "=SUM(A1:A9)", base + timedelta(hours=1), 5, "Ugx1"),
        Comment("luca", "risposta con\nnewline ed emoji 🎉", base + timedelta(hours=2), 1,
                "Ugx1.1", is_reply=True, parent_author="mario"),
        Comment("[utente rimosso]", "x" * 40_000, base + timedelta(hours=3), 0, "Ugx2"),
    ]
    return video, comments


class TestReportXlsx(unittest.TestCase):
    def _scrivi(self, write_only: bool, note: str = "") -> Path:
        video, comments = _dati_di_prova()
        tmp = Path(tempfile.mkdtemp())
        percorso = tmp / "report.xlsx"
        return write_report(
            percorso,
            video,
            comments,
            extracted_at=datetime(2026, 9, 23, 11, 45, tzinfo=timezone.utc),
            note=note,
            utc=True,
            write_only=write_only,
        )

    def test_struttura_e_stili(self) -> None:
        for write_only in (False, True):
            with self.subTest(write_only=write_only):
                salvato = self._scrivi(write_only)
                foglio = load_workbook(salvato)["Report"]

                self.assertEqual(foglio["A1"].value, "REPORT COMMENTI YOUTUBE")
                self.assertEqual(foglio["A3"].value, "Link video")
                self.assertEqual(foglio["B4"].value, "Canale di Prova")
                self.assertEqual(foglio["B5"].value, "Titolo di prova")
                self.assertEqual(foglio["B7"].value, 1_234_567)
                self.assertEqual(foglio["A12"].value, "COMMENTI")

                # 8 righe di metadati => intestazione della tabella commenti alla riga 13
                self.assertEqual(foglio["A13"].value, "Utente")
                self.assertEqual(foglio["F13"].value, "Risposta a")
                self.assertIsNone(foglio.freeze_panes)  # il foglio scorre tutto
                self.assertEqual(foglio.auto_filter.ref, "A13:F16")

                # Colori diversi tra etichetta e valore, e tra commento e risposta
                self.assertEqual(foglio["A3"].fill.fgColor.rgb[-6:], xlsx_report.BLUE_LABEL)
                self.assertEqual(foglio["B3"].fill.fgColor.rgb[-6:], xlsx_report.BLUE_VALUE)
                self.assertEqual(foglio["A13"].fill.fgColor.rgb[-6:], xlsx_report.NAVY)
                self.assertEqual(foglio["A15"].fill.fgColor.rgb[-6:], xlsx_report.REPLY_B)
                self.assertTrue(foglio["A13"].font.bold)

                # Il commento che inizia con "=" resta testo, non diventa una formula
                self.assertEqual(foglio["B14"].data_type, "s")
                self.assertEqual(foglio["B14"].value, "=SUM(A1:A9)")

                # Date come datetime reali, senza timezone, e numeri come numeri
                self.assertEqual(foglio["C14"].value, datetime(2024, 3, 14, 19, 30))
                self.assertEqual(foglio["D14"].value, 5)
                self.assertEqual(foglio["E15"].value, "Risposta")
                self.assertEqual(foglio["F15"].value, "mario")

                # Testo oltre il limite di Excel troncato senza errori di salvataggio
                self.assertTrue(foglio["B16"].value.endswith(TRUNCATE_MARK))

    def test_riga_note_sposta_la_tabella(self) -> None:
        salvato = self._scrivi(False, note="report parziale")
        foglio = load_workbook(salvato)["Report"]
        self.assertEqual(foglio["A11"].value, "Note")
        self.assertEqual(foglio["A14"].value, "Utente")
        self.assertIsNone(foglio.freeze_panes)

    def test_report_senza_commenti(self) -> None:
        video, _ = _dati_di_prova()
        tmp = Path(tempfile.mkdtemp())
        salvato = write_report(
            tmp / "vuoto.xlsx",
            video,
            [],
            extracted_at=datetime(2026, 9, 23, 11, 45, tzinfo=timezone.utc),
            note="I commenti sono disabilitati su questo video.",
        )
        foglio = load_workbook(salvato)["Report"]
        # 8 metadati + riga Note => intestazione alla riga 14, messaggio alla 15
        self.assertEqual(foglio["A15"].value, "Nessun commento estratto.")


if __name__ == "__main__":
    unittest.main(verbosity=2)
